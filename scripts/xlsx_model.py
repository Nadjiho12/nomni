#!/usr/bin/env python3
"""Spreadsheet modeling helper: build a formula-driven subscription forecast workbook, verify it by recalculating in LibreOffice,
and scan any .xlsx for formula errors. Requires openpyxl; recalculation requires LibreOffice (`soffice`).

  python xlsx_model.py build --out model.xlsx --new-per-month 20 --growth 0.05 --churn 0.05 --arpu 12 --gross-margin 0.85 --fixed 800 --cac 25 --cash 3000
  python xlsx_model.py verify model.xlsx          # recalculates, compares every Forecast cell with an independent Python model, checks for errors
  python xlsx_model.py check any.xlsx             # recalculates and lists #DIV/0!, #REF!, #NAME?, #VALUE!, #N/A, #NUM!, Err:### cells
The workbook separates inputs (Assumptions), calculations (Forecast, all formulas, no hard-coded numbers), and outputs (Summary with checks).
Exit codes: 0 ok; 1 mismatch or formula errors; 2 usage error; 3 LibreOffice or openpyxl unavailable.
"""
import argparse
import os
import shutil
import subprocess
import sys
import tempfile

ERRS = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
ASSUME = [("New customers in month 1", "new_per_month"), ("Monthly growth of new customers", "growth"), ("Monthly churn", "churn"), ("Revenue per customer per month (ARPU)", "arpu"),
          ("Gross margin", "gross_margin"), ("Fixed costs per month", "fixed"), ("Customer acquisition cost (CAC)", "cac"), ("Starting cash", "cash"), ("Starting customers", "start_customers")]


def die(msg, code=2):
    print("error: " + msg, file=sys.stderr)
    sys.exit(code)


def need_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        print("error: openpyxl is required: pip install openpyxl", file=sys.stderr)
        sys.exit(3)


def reference_model(p, months):
    """Independent Python model of the same forecast (mirrors scripts/biz_calc.py forecast)."""
    cust, cash, rows = p["start_customers"], p["cash"], []
    for m in range(1, months + 1):
        new = p["new_per_month"] * (1 + p["growth"]) ** (m - 1)
        lost = cust * p["churn"]
        cust = cust - lost + new
        mrr = cust * p["arpu"]
        gp = mrr * p["gross_margin"]
        spend = new * p["cac"]
        net = gp - spend - p["fixed"]
        cash += net
        rows.append(dict(month=m, new=new, churned=lost, customers=cust, mrr=mrr, gross_profit=gp, acq=spend, fixed=p["fixed"], net=net, cash=cash))
    return rows


def build(path, p, months):
    openpyxl = need_openpyxl()
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "Assumptions"
    a.append(["Assumption (inputs: blue cells are the only numbers to edit)", "Value"])
    for label, key in ASSUME:
        a.append([label, p[key]])
    for r in range(2, 2 + len(ASSUME)):
        a.cell(r, 2).font = Font(color="0000FF")
        a.cell(r, 2).fill = PatternFill("solid", fgColor="FFF2CC")
    for r in (3, 4, 6):
        a.cell(r, 2).number_format = "0.0%"
    a.column_dimensions["A"].width = 46
    f = wb.create_sheet("Forecast")
    f.append(["Month", "New customers", "Churned", "Customers (end)", "MRR", "Gross profit", "Acquisition spend", "Fixed costs", "Net cash flow", "Cash balance", "Break-even flag", "Cash-negative flag"])
    for i in range(months):
        r = i + 2
        prev_c = "Assumptions!$B$10" if i == 0 else f"D{r - 1}"
        prev_cash = "Assumptions!$B$9" if i == 0 else f"J{r - 1}"
        f.append([i + 1, f"=Assumptions!$B$2*(1+Assumptions!$B$3)^(A{r}-1)", f"={prev_c}*Assumptions!$B$4", f"={prev_c}-C{r}+B{r}", f"=D{r}*Assumptions!$B$5",
                  f"=E{r}*Assumptions!$B$6", f"=B{r}*Assumptions!$B$8", "=Assumptions!$B$7", f"=F{r}-G{r}-H{r}", f"={prev_cash}+I{r}", f"=IF(I{r}>=0,1,0)", f"=IF(J{r}<0,1,0)"])
    last = months + 1
    s = wb.create_sheet("Summary")
    s.append(["Output", "Value"])
    s.append(["Customers at end of horizon", f"=Forecast!D{last}"])
    s.append(["MRR at end of horizon", f"=Forecast!E{last}"])
    s.append(["Lowest cash balance (including starting cash)", f"=MIN(Assumptions!B9,Forecast!J2:J{last})"])
    s.append(["First month with non-negative net cash flow", f'=IFERROR(MATCH(1,Forecast!K2:K{last},0),"not within horizon")'])
    s.append(["First month cash is negative", f'=IFERROR(MATCH(1,Forecast!L2:L{last},0),"never")'])
    s.append(["CHECK: ending cash = starting cash + sum of net cash flow (should be 0)", f"=ROUND(Forecast!J{last}-(Assumptions!B9+SUM(Forecast!I2:I{last})),6)"])
    s.column_dimensions["A"].width = 64
    for col in "BCDEFGHIJKL":
        f.column_dimensions[col].width = 16
    wb.save(path)
    print(f"wrote {path} ({months} months, {months * 11 + 6} formula cells; verify with: xlsx_model.py verify {path})")


def recalc(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("error: LibreOffice (soffice) is required to recalculate formulas", file=sys.stderr)
        sys.exit(3)
    out = tempfile.mkdtemp(prefix="recalc-")
    prof = tempfile.mkdtemp(prefix="lo-profile-")
    r = subprocess.run([soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, path],
                       capture_output=True, text=True, timeout=240)
    res = os.path.join(out, os.path.basename(path))
    if not os.path.exists(res):
        die(f"LibreOffice could not recalculate {path}: {(r.stderr or r.stdout).strip()[:200]}", 1)
    shutil.rmtree(prof, ignore_errors=True)
    return res


def scan_errors(openpyxl, path):
    wb = openpyxl.load_workbook(path, data_only=True)
    bad, formulas = [], 0
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and (c.value in ERRS or c.value.startswith("Err:")):
                    bad.append(f"{ws.title}!{c.coordinate} = {c.value}")
    wf = openpyxl.load_workbook(path)
    for ws in wf:
        for row in ws.iter_rows():
            formulas += sum(1 for c in row if isinstance(c.value, str) and c.value.startswith("="))
    return bad, formulas


def cmd_build(a):
    if a.months < 1 or a.months > 120:
        die("--months must be between 1 and 120")
    p = dict(new_per_month=a.new_per_month, growth=a.growth, churn=a.churn, arpu=a.arpu, gross_margin=a.gross_margin, fixed=a.fixed, cac=a.cac, cash=a.cash, start_customers=a.start_customers)
    if not (0 <= a.churn <= 1) or not (0 <= a.gross_margin <= 1) or min(a.new_per_month, a.arpu, a.fixed, a.cac, a.start_customers) < 0:
        die("churn and gross margin must be in [0,1]; counts, prices, and costs must be >= 0")
    build(a.out, p, a.months)


def cmd_check(a):
    openpyxl = need_openpyxl()
    if not os.path.exists(a.file):
        die(f"{a.file} not found")
    res = recalc(os.path.abspath(a.file))
    bad, n = scan_errors(openpyxl, res)
    print(f"{n} formula cells recalculated; {len(bad)} error cell(s)")
    for b in bad[:50]:
        print("  ERROR", b)
    sys.exit(1 if bad else 0)


def cmd_verify(a):
    openpyxl = need_openpyxl()
    if not os.path.exists(a.file):
        die(f"{a.file} not found")
    wb0 = openpyxl.load_workbook(a.file)
    if "Assumptions" not in wb0.sheetnames or "Forecast" not in wb0.sheetnames:
        die("not a workbook built by this tool (needs Assumptions and Forecast sheets)")
    av = wb0["Assumptions"]
    p = {key: av.cell(2 + i, 2).value for i, (_, key) in enumerate(ASSUME)}
    months = wb0["Forecast"].max_row - 1
    ref = reference_model(p, months)
    res = recalc(os.path.abspath(a.file))
    bad, n = scan_errors(openpyxl, res)
    wb = openpyxl.load_workbook(res, data_only=True)
    ws = wb["Forecast"]
    cols = ["month", "new", "churned", "customers", "mrr", "gross_profit", "acq", "fixed", "net", "cash"]
    mism = []
    for i, row in enumerate(ref):
        for j, k in enumerate(cols):
            v = ws.cell(i + 2, j + 1).value
            if v is None or abs(float(v) - row[k]) > 1e-6 * max(1.0, abs(row[k])):
                mism.append(f"Forecast row {i + 2} {k}: sheet={v} python={row[k]:.6f}")
    chk = wb["Summary"].cell(7, 2).value
    print(f"{n} formulas, {len(bad)} error cells, {len(mism)} mismatches vs the independent Python model; reconciliation check cell = {chk}")
    for m in (bad + mism)[:20]:
        print("  PROBLEM", m)
    ok = not bad and not mism and chk == 0
    print("RESULT:", "workbook verified" if ok else "PROBLEMS FOUND")
    sys.exit(0 if ok else 1)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("--out", required=True); b.add_argument("--months", type=int, default=12)
    for name, default in (("new-per-month", 20.0), ("growth", 0.05), ("churn", 0.05), ("arpu", 12.0), ("gross-margin", 0.85), ("fixed", 800.0), ("cac", 25.0), ("cash", 3000.0), ("start-customers", 0.0)):
        b.add_argument(f"--{name}", type=float, default=default)
    b.set_defaults(fn=cmd_build)
    for n, fn in (("verify", cmd_verify), ("check", cmd_check)):
        p = sub.add_parser(n); p.add_argument("file"); p.set_defaults(fn=fn)
    a = ap.parse_args()
    a.fn(a)


if __name__ == "__main__":
    main()
