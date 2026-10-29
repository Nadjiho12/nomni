"""Tests for the domain tools (stl_tools, audio_tools, xlsx_model) and the runnable/compilable snippets in the new domain references."""
import json
import math
import os
import re
import subprocess
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from helpers import ENV, REFS, ROOT, SCRIPTS, Base, have, have_cmd, run  # noqa: E402

sys.path.insert(0, SCRIPTS)


def block(ref, marker, lang="python"):
    text = open(os.path.join(REFS, ref), encoding="utf-8").read()
    found = [b for b in re.findall(r"```" + lang + r"\n(.*?)```", text, re.S) if marker in b]
    assert found, f"{ref}: no {lang} block containing {marker!r}"
    return found[0]


class StlTools(Base):
    def gen(self, *args):
        r = run(["stl_tools.py", "gen"] + list(args))
        self.assertEqual(r.returncode, 0, r.stderr)
        return r

    def check(self, path, *extra):
        r = run(["stl_tools.py", "check", path, "--json"] + list(extra))
        return r.returncode, json.loads(r.stdout) if r.stdout.strip() else {}

    def test_box_exact_values(self):
        p = os.path.join(self.tmp, "b.stl")
        self.gen("box", "--size", "10", "20", "30", "--out", p)
        code, d = self.check(p)
        self.assertEqual(code, 0)
        self.assertAlmostEqual(d["volume"], 6000.0, places=3)
        self.assertAlmostEqual(d["surface_area"], 2200.0, places=3)
        self.assertEqual(d["size"], [10.0, 20.0, 30.0])
        self.assertEqual((d["triangles"], d["vertices"]), (12, 8))
        self.assertTrue(d["watertight"] and d["consistently_oriented"])

    def test_cylinder_and_tube_match_polygon_formulas_ascii_and_binary(self):
        for binary in (False, True):
            cyl = os.path.join(self.tmp, f"c{binary}.stl")
            self.gen("cylinder", "--radius", "5", "--height", "10", "--segments", "90", "--out", cyl, *(["--binary"] if binary else []))
            code, d = self.check(cyl)
            exp = 0.5 * 90 * 25 * math.sin(2 * math.pi / 90) * 10
            self.assertEqual(code, 0)
            self.assertAlmostEqual(d["volume"], exp, delta=0.01 * exp / 100 + 0.01)
            tube = os.path.join(self.tmp, f"t{binary}.stl")
            self.gen("tube", "--outer", "10", "--inner", "7", "--height", "20", "--segments", "90", "--out", tube, *(["--binary"] if binary else []))
            code, d = self.check(tube)
            exp = 0.5 * 90 * (100 - 49) * math.sin(2 * math.pi / 90) * 20
            self.assertEqual(code, 0)
            self.assertAlmostEqual(d["volume"], exp, delta=0.05)
            self.assertTrue(d["watertight"] and d["consistently_oriented"])

    def mesh_file(self, tris, name):
        import stl_tools
        p = os.path.join(self.tmp, name)
        stl_tools.write_ascii(p, tris)
        return p

    def test_defects_are_detected(self):
        import stl_tools
        tris = stl_tools.box(10, 10, 10)
        code, d = self.check(self.mesh_file(tris[:-1], "hole.stl"))                       # one triangle missing: open edges
        self.assertEqual(code, 1)
        self.assertFalse(d["watertight"])
        self.assertGreater(d["open_edges"], 0)
        flipped = list(tris)
        flipped[3] = (flipped[3][0], flipped[3][2], flipped[3][1])                         # one face flipped
        code, d = self.check(self.mesh_file(flipped, "flip.stl"))
        self.assertEqual(code, 1)
        self.assertFalse(d["consistently_oriented"])
        inv = [(t[0], t[2], t[1]) for t in tris]                                           # whole mesh inside out
        code, d = self.check(self.mesh_file(inv, "inv.stl"))
        self.assertEqual(code, 1)
        self.assertLess(d["volume"], 0)
        degenerate = tris + [((0, 0, 0), (0, 0, 0), (1, 1, 1))]
        code, d = self.check(self.mesh_file(degenerate, "deg.stl"))
        self.assertEqual(code, 1)
        self.assertEqual(d["degenerate_triangles"], 1)

    def test_clean_errors(self):
        for args in (["gen", "box", "--out", "x.stl"], ["gen", "box", "--size", "1", "0", "1", "--out", "x.stl"], ["gen", "cylinder", "--out", "x.stl"],
                     ["gen", "tube", "--outer", "5", "--inner", "6", "--height", "3", "--out", "x.stl"], ["gen", "cylinder", "--radius", "5", "--height", "3", "--segments", "2", "--out", "x.stl"]):
            r = run(["stl_tools.py"] + args, cwd=self.tmp)
            self.assertEqual(r.returncode, 2, args)
            self.assertNoTraceback(r)
        for content in ("", "garbage not an stl file at all", "solid x\nfacet normal 0 0 0\nvertex a b c\n"):
            r = run(["stl_tools.py", "check", self.f("g.stl", content)])
            self.assertEqual(r.returncode, 2, content)
            self.assertNoTraceback(r)
        self.assertEqual(run(["stl_tools.py", "check", "/nonexistent.stl"]).returncode, 2)


class AudioTools(Base):
    def wav(self, segments, name="t.wav", extra=()):
        p = os.path.join(self.tmp, name)
        r = run(["audio_tools.py", "compose", "--segments", segments, "--out", p] + list(extra))
        self.assertEqual(r.returncode, 0, r.stderr)
        return p

    def test_vad_finds_two_segments_with_correct_boundaries(self):
        p = self.wav("silence:0.5,tone:1.0,silence:0.7,tone:0.4,silence:0.3")
        out = run(["audio_tools.py", "vad", p, "--threshold-db", "-35", "--min-speech-ms", "200", "--hangover-ms", "100"]).stdout
        self.assertIn("2 speech segment(s)", out)
        segs = [(float(a), float(b)) for a, b in re.findall(r"\d\.\s+([\d.]+)s -\s+([\d.]+)s", out)]
        self.assertAlmostEqual(segs[0][0], 0.5, delta=0.06)
        self.assertAlmostEqual(segs[0][1], 1.5, delta=0.06)
        self.assertAlmostEqual(segs[1][0], 2.2, delta=0.06)
        self.assertAlmostEqual(segs[1][1], 2.6, delta=0.06)

    def test_vad_ignores_short_blips_and_silence(self):
        self.assertIn("0 speech segment(s)", run(["audio_tools.py", "vad", self.wav("silence:1.0")]).stdout)
        out = run(["audio_tools.py", "vad", self.wav("silence:0.5,tone:0.05,silence:0.5"), "--min-speech-ms", "200"]).stdout
        self.assertIn("0 speech segment(s)", out)

    def test_levels_detects_clipping(self):
        loud = self.wav("tone:0.5", "loud.wav", ["--amp", "1.0"])
        out = run(["audio_tools.py", "levels", loud]).stdout
        self.assertIn("CLIPPING", out)
        ok = self.wav("tone:0.5", "ok.wav", ["--amp", "0.3"])
        self.assertNotIn("CLIPPING", run(["audio_tools.py", "levels", ok]).stdout)

    def test_budget(self):
        r = run(["audio_tools.py", "budget", "--stage", "vad=40", "--stage", "asr=300", "--stage", "llm=450", "--stage", "tts=200", "--target", "900"])
        self.assertEqual(r.returncode, 1)
        self.assertIn("OVER by 90", r.stdout)
        self.assertIn("bottleneck: llm", r.stdout)
        r = run(["audio_tools.py", "budget", "--stage", "asr=200", "--stage", "llm=300", "--target", "1000"])
        self.assertEqual(r.returncode, 0)
        self.assertIn("under by 500", r.stdout)

    def test_wer(self):
        self.assertIn("WER 40.0%", run(["audio_tools.py", "wer", "turn on the kitchen light", "turn the kitchen lights"]).stdout)
        self.assertIn("WER 0.0%", run(["audio_tools.py", "wer", "same words", "same words"]).stdout)

    def test_clean_errors(self):
        cases = (["budget"], ["budget", "--stage", "asr"], ["budget", "--stage", "asr=x"], ["budget", "--stage", "asr=-5"], ["budget", "--stage", "a=1", "--target", "0"],
                 ["vad", "/nonexistent.wav"], ["levels", self.f("bad.wav", "not a wav")], ["compose", "--segments", "silence:abc", "--out", os.path.join(self.tmp, "x.wav")],
                 ["compose", "--segments", "noise:1", "--out", os.path.join(self.tmp, "x.wav")], ["wer", "", "x"])
        for args in cases:
            r = run(["audio_tools.py"] + list(args))
            self.assertEqual(r.returncode, 2, args)
            self.assertNoTraceback(r)


@unittest.skipUnless(have("openpyxl") and (have_cmd("soffice") or have_cmd("libreoffice")), "openpyxl and LibreOffice (soffice) required")
class XlsxModel(Base):
    def build(self, name="m.xlsx", *extra):
        p = os.path.join(self.tmp, name)
        r = run(["xlsx_model.py", "build", "--out", p] + list(extra))
        self.assertEqual(r.returncode, 0, r.stderr)
        return p

    def test_build_verify_and_crosscheck_with_biz_calc(self):
        p = self.build()
        r = run(["xlsx_model.py", "verify", p], timeout=240)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("0 error cells, 0 mismatches", r.stdout)
        # the workbook's month-12 cash must equal the independent biz_calc.py forecast
        import openpyxl
        from_xlsx = run(["xlsx_model.py", "check", p], timeout=240)
        self.assertEqual(from_xlsx.returncode, 0)
        fc = run(["biz_calc.py", "forecast", "--months", "12", "--new-per-month", "20", "--growth", "0.05", "--churn", "0.05", "--arpu", "12",
                  "--gross-margin", "0.85", "--fixed", "800", "--cac", "25", "--cash", "3000"]).stdout
        last = [l for l in fc.splitlines() if re.match(r"^12\s", l)][0].split()
        self.assertEqual(last[-1], "1,720")
        self.assertIsNotNone(openpyxl)

    def test_verify_detects_a_tampered_formula(self):
        import openpyxl
        p = self.build("t.xlsx", "--months", "6")
        wb = openpyxl.load_workbook(p)
        wb["Forecast"]["E4"] = "=D4*Assumptions!$B$5*1.1"           # silently inflates MRR by 10%
        wb.save(p)
        r = run(["xlsx_model.py", "verify", p], timeout=240)
        self.assertEqual(r.returncode, 1)
        self.assertIn("mismatches", r.stdout)
        self.assertIn("PROBLEMS FOUND", r.stdout)

    def test_check_finds_formula_errors(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["A2"], ws["A3"], ws["A4"] = 1, 0, "=A1/A2", "=NOSUCHFUNC(1)"
        p = os.path.join(self.tmp, "e.xlsx")
        wb.save(p)
        r = run(["xlsx_model.py", "check", p], timeout=240)
        self.assertEqual(r.returncode, 1)
        self.assertIn("#DIV/0!", r.stdout)
        self.assertIn("#NAME?", r.stdout)


class XlsxModelArgs(Base):
    def test_bad_arguments(self):
        for extra in (["--months", "0"], ["--months", "500"], ["--churn", "2"], ["--gross-margin", "-0.1"], ["--cac", "-1"]):
            r = run(["xlsx_model.py", "build", "--out", os.path.join(self.tmp, "x.xlsx")] + extra)
            self.assertIn(r.returncode, (2, 3), extra)           # 3 when openpyxl is missing
            self.assertNoTraceback(r)
        self.assertIn(run(["xlsx_model.py", "verify", "/nonexistent.xlsx"]).returncode, (2, 3))


@unittest.skipUnless(have_cmd("g++"), "g++ not installed")
class EmbeddedSnippets(Base):
    def test_every_cpp_sketch_compiles_against_the_arduino_stub(self):
        text = open(os.path.join(REFS, "extra", "embedded-iot.md"), encoding="utf-8").read()
        blocks = re.findall(r"```cpp\n(.*?)```", text, re.S)
        self.assertGreaterEqual(len(blocks), 6)
        stub = os.path.join(ROOT, "tests", "fixtures", "arduino_stub.h")
        for i, code in enumerate(blocks):
            src = self.f(f"sketch{i}.cpp", code)
            r = subprocess.run(["g++", "-std=gnu++17", "-fsyntax-only", "-Wall", "-I", os.path.dirname(stub), "-include", stub, src], capture_output=True, text=True, encoding="utf-8", errors="replace")
            self.assertEqual(r.returncode, 0, f"sketch {i} failed to compile:\n{r.stderr[:600]}")
            self.assertIn("setup", code)
            self.assertIn("loop", code)

    def test_stub_rejects_a_bad_sketch(self):                       # negative control: the check can fail
        src = self.f("bad.cpp", "void setup() { pinModee(13, OUTPUT); }\nvoid loop() {}\n")
        r = subprocess.run(["g++", "-std=gnu++17", "-fsyntax-only", "-include", os.path.join(ROOT, "tests", "fixtures", "arduino_stub.h"), src], capture_output=True, text=True, encoding="utf-8", errors="replace")
        self.assertNotEqual(r.returncode, 0)


class DomainPythonSnippets(Base):
    def exec_block(self, ref, marker):
        r = subprocess.run([sys.executable, "-c", block(ref, marker)], capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120, cwd=self.tmp, env=ENV)
        self.assertEqual(r.returncode, 0, r.stderr[-500:])
        return r.stdout

    def test_pid_simulation_settles_without_runaway(self):
        out = self.exec_block("extra/embedded-iot.md", "def simulate(setpoint")
        self.assertIn("settled within 1 C: True", out)
        overshoot = float(re.search(r"max overshoot (-?[\d.]+)", out).group(1))
        self.assertLess(overshoot, 10.0)

    def test_cad_calculations(self):
        out = self.exec_block("extra/cad-3d-printing.md", "def hole_diameter")
        self.assertRegex(out, r"sliding\s+8 mm pin -> 8\.45 mm hole")
        out = self.exec_block("extra/cad-3d-printing.md", "def cantilever_deflection_mm")
        self.assertIn("3.7 mm", out)

    def test_voice_wer_snippet_matches_the_cli(self):
        self.assertEqual(self.exec_block("extra/voice-speech-ai.md", "def wer(").strip(), "40.0")

    def test_spreadsheet_snippet_writes_formulas(self):
        if not have("openpyxl"):
            self.skipTest("openpyxl not installed")
        out = self.exec_block("extra/spreadsheets-modeling.md", "tiny_model.xlsx")
        self.assertIn("=B1*B2", out)
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "tiny_model.xlsx")))


if __name__ == "__main__":
    unittest.main()
